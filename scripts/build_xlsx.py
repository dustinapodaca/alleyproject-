import re, datetime as dt, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
import sys
sys.path.insert(0, '/home/claude/plan/paint')
from idmap import KIT_COLORS, HERO_COLORS, VISIBLE

SRC = "/root/.claude/uploads/d4a4280e-07db-56a9-bd1c-34163b3262d6/c3030a96-KingArthur_Copy.xlsx"
PLAN = "/home/claude/mill-alley-tracker/scripts/plan.md"
OUT = "/home/claude/plan/Derelict-Mill-Alley-Tracker.xlsx"
s = open(PLAN).read()

def table_after(label):
    i = s.index(label); seg = s[i:].split("\n\n", 2)[1]
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in seg.strip().split("\n")]
    return [dict(zip(rows[0], r)) for r in rows[2:]]

# Template colours
HEAD, BAND = "FF76A5AF", "FFA2C4C9"
DONE, WIP, TODO, SEP = "FF93C47D", "FFFFD966", "FFE06666", "FFCCCCCC"
fill = lambda c: PatternFill("solid", fgColor=c)
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)

wb = openpyxl.load_workbook(SRC)
if "Sheet2" in wb.sheetnames:
    del wb["Sheet2"]  # unrelated reading list from the original file
ws = wb["Sheet1"]; ws.title = "Mill Alley"
# clear template content and fills, keep sheet-level settings
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=17):
    for c in row:
        c.value = None; c.fill = PatternFill(); c.comment = None

ws["A1"] = "Derelict Mill Alley · Assets"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
ws["I1"] = "Scheme · Oct 19, 2026 – Jan 31, 2027 (Wed/Thu off)"; ws["I1"].font = Font(name="Arial", size=14, bold=True)

# ---------- Left: asset tracker ----------
kitA, kitB, kitC, kitD = (table_after(f"**Family {x}:") for x in "ABCD")
heroes = table(s) if False else None
i = s.index("## Unique hero features"); seg = s[i:].split("\n\n")[2]
rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in seg.strip().split("\n")]
heroes = [dict(zip(rows[0], r)) for r in rows[2:]]
props = table_after("**Props and debris**")
foliage = table_after("**Foliage**\n")
decals = table_after("**Decal sets**")
mats = table_after("**Tiling materials")
trims = [("T01", "Stone trim sheet", "2048 px, 10.24 px/cm"), ("T02", "Metal trim sheet", "2048 px, 10.24 px/cm"), ("T03", "Timber trim sheet", "2048 px, 10.24 px/cm")]
HERO_SIZES = {"H1": "500 span x 340 h x 250 d", "H2": "180 x 2800 h x 180", "H3": "120 x 200 drop x 380", "H6": "240 x 320", "H7": "140 x 220"}

groups = [
    ("Kit A · Left mill", ["Blockout", "Model", "UVs", "Material", "In UE5"], [(r["ID"], r["Piece"], r["Size (cm)"]) for r in kitA]),
    ("Kit B · Right building", ["Blockout", "Model", "UVs", "Material", "In UE5"], [(r["ID"], r["Piece"], r["Size (cm)"]) for r in kitB]),
    ("Kit C · Structure", ["Blockout", "Model", "UVs", "Material", "In UE5"], [(r["ID"], r["Piece"], r["Size (cm)"]) for r in kitC]),
    ("Kit D · Ground", ["Blockout", "Model", "UVs", "Material", "In UE5"], [(r["ID"], r["Piece"], r["Size (cm)"]) for r in kitD]),
    ("Hero features", ["Blockout", "Highpoly", "Lowpoly", "Texture", "In UE5"], [(r["#"], r["Feature"], HERO_SIZES.get(r["#"], "")) for r in heroes]),
    ("Props", ["Blockout", "Highpoly", "Lowpoly", "Texture", "In UE5"], [(r["ID"], r["Asset"], r["Size (cm)"]) for r in props]),
    ("Vegetation", ["Model", "Atlas", "Material", "Wind", "Placed"], [(r["ID"], r["Asset"], "") for r in foliage]),
    ("Tile-Textures", ["Source", "Customize", "Masks", "Bake", "In UE5"], [(r["ID"], r["Material"], r["Tile size (cm)"] + " tile") for r in mats]),
    ("Trim sheets", ["Model", "Bake", "Texture", "Material", "In UE5"], trims),
    ("Decals", ["Source", "Bake", "Texture", "Material", "Placed"], [(r["ID"], r["Decal set"], "") for r in decals]),
]
r = 2
for name, stages, items in groups:
    for col in "ABCDEFG":
        ws[f"{col}{r}"].fill = fill(HEAD); ws[f"{col}{r}"].font = F
        ws[f"{col}{r+1}"].fill = fill(BAND); ws[f"{col}{r+1}"].font = F
    ws[f"A{r}"] = name
    ws[f"B{r}"] = "Size (cm)"
    ws[f"H{r}"] = "ID"; ws[f"H{r}"].fill = fill(HEAD); ws[f"H{r}"].font = Font(name="Arial", size=8)
    ws[f"H{r+1}"].fill = fill(BAND)
    for col, st in zip("CDEFG", stages):
        ws[f"{col}{r}"] = st
    r += 2
    for iid, label, size in items:
        ws[f"A{r}"] = f"{iid} {label}"; ws[f"A{r}"].font = F
        ws[f"B{r}"] = size or None; ws[f"B{r}"].font = Font(name="Arial", size=9, color="FF555555")
        for col in "CDEFG":
            ws[f"{col}{r}"].fill = fill(TODO)
        hexc = KIT_COLORS.get(iid) or HERO_COLORS.get(iid)
        if hexc:
            ws[f"H{r}"].fill = fill("FF" + hexc.lstrip("#").upper())
            ws[f"H{r}"].comment = Comment(f"ID map colour {hexc}" + ("" if iid in VISIBLE or iid in HERO_COLORS else " (not visible in the reference photo)"), "Mill Alley plan")
        r += 1
r += 2
for colour, label in [(DONE, "Done"), (WIP, "In Progress"), (TODO, "Not Started")]:
    ws[f"A{r}"].fill = fill(colour); ws[f"B{r}"] = label; ws[f"B{r}"].font = F; r += 1
ws[f"A{r+1}"] = "Change a stage cell's fill colour to update its status. Column H is the ID-map colour from the paint-over (see Paint-over sheet)."; ws[f"A{r+1}"].font = Font(name="Arial", size=9, italic=True, color="FF666666")

# ---------- Right: weekly schedule ----------
start = dt.date(2026, 10, 19)
weeks = list(re.finditer(r"^### Week (\d+) · (.+?) · (.+)$\n\n((?:- \[ \].*\n?)+)", s, flags=re.M))
assert len(weeks) == 15

def short(t):
    t = t.split("*Done when:*")[0].strip()
    cut = len(t)
    for sep in [":", ";", " ("]:
        k = t.find(sep)
        if 8 < k < cut: cut = k
    out = t[:cut].strip().rstrip(".")
    return out.replace("`", "")

cols = "IJKLMNO"
r = 2
for m in weeks:
    n = int(m.group(1)); mon = start + dt.timedelta(days=7 * (n - 1))
    days = {}
    for line in m.group(4).strip().split("\n"):
        lab, txt = re.match(r"- \[ \] \*\*(.+?)\*\* · (.*)$", line).groups()
        days[lab] = txt
    for col in cols:
        ws[f"{col}{r}"].fill = fill(HEAD); ws[f"{col}{r}"].font = F
        ws[f"{col}{r+1}"].fill = fill(BAND)
        ws[f"{col}{r+3}"].fill = fill(SEP)
    ws[f"I{r}"] = f"Week {n} · {m.group(3)}"
    for k, col in enumerate(cols):
        d = mon + dt.timedelta(days=k)
        c = ws[f"{col}{r+1}"]; c.value = d; c.number_format = "ddd m/d"; c.font = F
        c.alignment = Alignment(horizontal="center")
        t = ws[f"{col}{r+2}"]
        key = f"{d:%a %b} {d.day}"
        if d.weekday() in (2, 3):
            t.value = "Day off"; t.font = Font(name="Arial", size=10, italic=True, color="FF999999")
        else:
            full = days[key]
            lab = short(full)
            rv = re.search(r"Review (\d+)", full)
            if rv and "Review" not in lab:
                lab += f" + Review {rv.group(1)}"
            t.value = lab
            buffer = full.startswith("Buffer")
            t.font = Font(name="Arial", size=9, bold=not buffer, italic=buffer, color="FF666666" if buffer else None)
            if not buffer:
                body = full.replace("`", "").replace("*Done when:*", "\n\nDone when:")
                cm = Comment(body, "Mill Alley plan"); cm.width = 320; cm.height = 160
                t.comment = cm
        t.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.row_dimensions[r + 2].height = 58
    r += 4

for col in cols: ws.column_dimensions[col].width = 22
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 26
ws.column_dimensions["H"].width = 4.5
ws.freeze_panes = "A2"

ms = wb.create_sheet("Measurements")
ms["A1"] = "Reference breakdown · 1:1 scale"; ms["A1"].font = Font(name="Arial", size=14, bold=True)
ms["A2"] = "Estimates from human-scale anchors (100 cm handrail, 18 cm risers, 210-260 cm doors, Victorian brick coursing). Confirm by camera-matching the photo in Blender."
ms["A2"].font = Font(name="Arial", size=9, italic=True, color="FF666666")
meas = table_after("**Measurements**")
hdr = ["Element", "Width (cm)", "Height (cm)", "Depth (cm)", "Notes"]
for j, h in enumerate(hdr):
    c = ms.cell(row=4, column=j + 1, value=h); c.fill = fill(HEAD); c.font = F
for i, row in enumerate(meas):
    for j, h in enumerate(hdr):
        v = row[h]
        try:
            v = float(v) if "." in v else int(v)
        except ValueError:
            v = v or None
        c = ms.cell(row=5 + i, column=j + 1, value=v); c.font = F
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    if i % 2: 
        for j in range(5): ms.cell(row=5 + i, column=j + 1).fill = fill("FFEEF3F4")
r2 = 5 + len(meas) + 1
ms.cell(row=r2, column=1, value="Metrics to hold").font = FB
metric_lines = [l[2:] for l in s.split("**Metrics to hold for the whole project**")[1].split("\n\n")[1].split("\n") if l.startswith("- ")]
for k, line in enumerate(metric_lines):
    c = ms.cell(row=r2 + 1 + k, column=1, value=line); c.font = F
for col, w in zip("ABCDE", [30, 14, 14, 18, 52]):
    ms.column_dimensions[col].width = w
ms.freeze_panes = "A5"

po = wb.create_sheet("Paint-over")
po["A1"] = "Modular paint-over · ID map colours"; po["A1"].font = Font(name="Arial", size=14, bold=True)
po["A2"] = "Same colours as column H on the Mill Alley sheet. Use them for ID maps, kit sheets and the blockout."
po["A2"].font = Font(name="Arial", size=9, italic=True, color="FF666666")
img = XLImage("/home/claude/plan/paint/mill-alley-paintover.png")
scale = 1100 / img.width; img.width = int(img.width * scale); img.height = int(img.height * scale)
po.add_image(img, "F4")
hdrs = ["ID", "Piece", "Colour", "Hex", "In photo"]
for j, h in enumerate(hdrs):
    c = po.cell(row=4, column=j + 1, value=h); c.fill = fill(HEAD); c.font = F
names = {}
for _, _, items in groups:
    for it in items: names[it[0]] = it[1]
rr = 5
for pid, hexc in list(KIT_COLORS.items()) + list(HERO_COLORS.items()):
    po.cell(row=rr, column=1, value=pid).font = F
    po.cell(row=rr, column=2, value=names.get(pid, "")).font = F
    po.cell(row=rr, column=3).fill = fill("FF" + hexc.lstrip("#").upper())
    po.cell(row=rr, column=4, value=hexc.upper()).font = Font(name="Arial", size=9, color="FF555555")
    seen = pid in VISIBLE or pid in HERO_COLORS or pid == "C05"
    po.cell(row=rr, column=5, value="Yes" if seen else "—").font = F
    rr += 1
for col, w in zip("ABCDE", [6, 30, 8, 10, 9]):
    po.column_dimensions[col].width = w
wb.save(OUT)
print("saved", OUT, "rows", ws.max_row)
